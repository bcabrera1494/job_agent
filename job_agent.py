#!/usr/bin/env python3
"""
Daily remote-job finder + application tailor.

What it does each run:
  1. Pulls open jobs from public job-board APIs (Greenhouse, Lever, Ashby, Remotive).
  2. Keeps only remote roles that match your title filters, are open to California,
     don't need heavy travel, and pay at least your salary floor.
  3. Skips anything it has already seen on a previous day.
  4. Scores fit with an LLM, then tailors a resume + cover letter for the best few.
  5. Writes output/<date>/summary.md (table), summary.csv, and one file per job.

It never applies for you. You review everything before sending.

Usage:
  python job_agent.py              # full run
  python job_agent.py --dry-run    # fetch + filter only, no LLM calls
  python job_agent.py --selftest   # offline tests of the parsing/filter logic
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import html
import json
import logging
import os
import pathlib
import re
import sqlite3
import sys
import time
from dataclasses import dataclass, field, asdict

import requests
import yaml

HERE = pathlib.Path(__file__).parent
UA = {"User-Agent": "personal-job-search-agent/1.0 (+personal use)"}
log = logging.getLogger("job_agent")


# --------------------------------------------------------------------------- #
# Data model
# --------------------------------------------------------------------------- #
@dataclass
class Job:
    source: str
    company: str
    title: str
    location: str
    url: str
    description: str
    salary_min: int | None = None
    salary_max: int | None = None
    salary_source: str = ""          # "api" | "text" | "text-single" | ""
    remote: bool = False
    flags: list[str] = field(default_factory=list)
    fit_score: int | None = None
    fit_verdict: str = ""
    fit_why: list[str] = field(default_factory=list)
    fit_gaps: list[str] = field(default_factory=list)

    @property
    def key(self) -> str:
        raw = f"{self.company}|{self.title}|{self.url}".lower()
        return hashlib.sha1(raw.encode()).hexdigest()

    @property
    def salary_display(self) -> str:
        if self.salary_min is None:
            return "Not listed"
        if self.salary_min == self.salary_max:
            return f"${self.salary_min:,}"
        return f"${self.salary_min:,} - ${self.salary_max:,}"


# --------------------------------------------------------------------------- #
# Text helpers
# --------------------------------------------------------------------------- #
def strip_html(raw: str) -> str:
    raw = html.unescape(raw or "")
    raw = re.sub(r"(?i)<br\s*/?>|</p>|</li>|</h\d>", "\n", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    raw = html.unescape(raw)
    raw = re.sub(r"[ \t\r\f\v]+", " ", raw)
    raw = re.sub(r"\n\s*\n+", "\n\n", raw)
    return raw.strip()


_AMOUNT = re.compile(
    r"(?<![A-Za-z])\$\s?(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)(?!\d)\s?([kK])?"
    r"(?!\s?(?:[mMbB]\b|million|billion|%))"
)
_NUMPAT = r"(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)"
_RANGE = re.compile(
    r"(?<![A-Za-z])\$\s?" + _NUMPAT + r"\s?([kK])?\s*(?:-|\u2013|\u2014|to)\s*\$?\s?" + _NUMPAT + r"(?!\d)\s?([kK])?"
)
_HOURLY = re.compile(r"/\s?h(?:ou)?r|per\s+hour|an\s+hour|hourly", re.I)
_SALARY_CUE = re.compile(r"salary|base pay|pay range|compensation|pay:|range", re.I)


def parse_salary(text: str) -> tuple[int | None, int | None, str]:
    """Return (min_annual, max_annual, kind). Handles $80,000-$95,000, $70K-$80K,
    'USD 90,000 - 130,000', and hourly rates (x2080). Ignores '$35M', '401(k)', CA$."""
    if not text:
        return None, None, ""
    text = re.sub(r"\b(?:USD|US\$)\s?(?=\d)", "$", text)

    def amounts(chunk: str):
        found = []
        for m in _AMOUNT.finditer(chunk):
            val = float(m.group(1).replace(",", ""))
            ctx = chunk[m.end(): m.end() + 16]
            if m.group(2):
                val *= 1000
            elif _HOURLY.search(ctx):
                val *= 2080
            elif val < 1000:
                continue
            if 20_000 <= val <= 500_000:
                found.append((m.start(), int(round(val))))
        return found

    # Prefer text right after a salary cue ("salary range", "compensation", ...)
    windows = [text[m.start(): m.start() + 300] for m in _SALARY_CUE.finditer(text)]
    def range_in(chunk: str):
        for m in _RANGE.finditer(chunk):
            k1, k2 = m.group(2), m.group(4)
            v1, v2 = float(m.group(1).replace(",", "")), float(m.group(3).replace(",", ""))
            if k1 and not k2 and v2 < 1000:
                k2 = k1                      # "$70-80K"
            ctx = chunk[m.end(): m.end() + 16]
            mult = 1000 if (k1 or k2) else (2080 if _HOURLY.search(ctx) else 1)
            if mult == 1 and v1 < 1000:
                continue
            lo, hi = v1 * (1000 if k1 else mult), v2 * (1000 if k2 else mult)
            if 20_000 <= min(lo, hi) and max(lo, hi) <= 500_000:
                return int(min(lo, hi)), int(max(lo, hi))
        return None

    for chunk in windows + [text]:
        r = range_in(chunk)
        if r:
            return r[0], r[1], "text"
        found = amounts(chunk)
        for (p1, v1), (p2, v2) in zip(found, found[1:]):
            if p2 - p1 < 60:
                return min(v1, v2), max(v1, v2), "text"
        if len(found) == 1 and chunk is not text:
            return found[0][1], found[0][1], "text-single"
    found = amounts(text)
    if len(found) == 1:
        return found[0][1], found[0][1], "text-single"
    return None, None, ""


# --------------------------------------------------------------------------- #
# Filters
# --------------------------------------------------------------------------- #
US_STATES = [
    "alabama", "alaska", "arizona", "arkansas", "california", "colorado", "connecticut",
    "delaware", "florida", "georgia", "hawaii", "idaho", "illinois", "indiana", "iowa",
    "kansas", "kentucky", "louisiana", "maine", "maryland", "massachusetts", "michigan",
    "minnesota", "mississippi", "missouri", "montana", "nebraska", "nevada",
    "new hampshire", "new jersey", "new mexico", "new york", "north carolina",
    "north dakota", "ohio", "oklahoma", "oregon", "pennsylvania", "rhode island",
    "south carolina", "south dakota", "tennessee", "texas", "utah", "vermont",
    "virginia", "washington", "west virginia", "wisconsin", "wyoming",
]
US_TOKEN = re.compile(r"\b(united states|usa|u\.s\.|us|america)\b", re.I)


def location_eligibility(location: str, cfg: dict) -> str:
    """'ok' (US/California explicitly fine), 'check' (just 'Remote'), or 'no'."""
    loc = (location or "").lower()
    home = cfg["profile"].get("home_state", "California").lower()
    foreign = [t.lower() for t in cfg["filters"].get("exclude_location_tokens", [])]
    has_us = bool(US_TOKEN.search(loc))
    if any(re.search(rf"\b{re.escape(t)}\b", loc) for t in foreign) and not has_us:
        return "no"
    states = [s for s in US_STATES if re.search(rf"\b{s}\b", loc)]
    if states and home not in states and not has_us:
        return "no"
    if home in states or has_us:
        return "ok"
    return "check"


_TRAVEL = [
    re.compile(r"travel[^.\n]{0,80}?(\d{1,2})\s?%", re.I),
    re.compile(r"(\d{1,2})\s?%[^.\n]{0,40}?travel", re.I),
    re.compile(r"up to\s+(\d{1,2})\s?%\s+travel", re.I),
]


def max_travel_percent(text: str) -> int | None:
    vals = [int(m.group(1)) for rx in _TRAVEL for m in rx.finditer(text or "")]
    return max(vals) if vals else None


def title_ok(title: str, cfg: dict) -> bool:
    f = cfg["filters"]
    inc = [re.compile(p, re.I) for p in f.get("title_include", [])]
    exc = [re.compile(p, re.I) for p in f.get("title_exclude", [])]
    return any(r.search(title) for r in inc) and not any(r.search(title) for r in exc)


def domain_ok(job: Job, cfg: dict) -> bool:
    kws = [k.lower() for k in cfg["filters"].get("domain_keywords", [])]
    blob = f"{job.company} {job.title} {job.description[:3000]}".lower()
    return any(k in blob for k in kws)


def apply_filters(job: Job, cfg: dict, trusted: bool) -> tuple[bool, str]:
    p = cfg["profile"]
    if not job.remote:
        return False, "not remote"
    if not title_ok(job.title, cfg):
        return False, "title"
    if not trusted and not domain_ok(job, cfg):
        return False, "not education/nonprofit"
    elig = location_eligibility(job.location, cfg)
    if elig == "no":
        return False, "location"
    if elig == "check":
        job.flags.append("Confirm US/California eligibility")
    travel = max_travel_percent(job.description)
    if travel is not None and travel > p.get("max_travel_percent", 25):
        return False, f"travel {travel}%"
    if travel:
        job.flags.append(f"Travel up to {travel}%")
    if job.salary_min is None:
        if not p.get("allow_unknown_salary", False):
            return False, "salary unknown"
        job.flags.append("Salary not listed")
    else:
        basis = job.salary_max if p.get("salary_basis", "max") == "max" else job.salary_min
        if basis < p["min_salary"]:
            return False, "salary below floor"
        if job.salary_source == "text-single":
            job.flags.append("Salary parsed from text: verify")
    return True, ""


# --------------------------------------------------------------------------- #
# Sources (public job-board APIs; no HTML scraping)
# --------------------------------------------------------------------------- #
def get_json(url: str, params: dict | None = None):
    for attempt in range(3):
        try:
            r = requests.get(url, params=params, headers=UA, timeout=30)
            if r.status_code == 404:
                log.warning("404 (check the board slug): %s", url)
                return None
            r.raise_for_status()
            return r.json()
        except requests.RequestException as e:
            log.warning("fetch failed (%s) attempt %d: %s", e, attempt + 1, url)
            time.sleep(2 * (attempt + 1))
    return None


def _annual(value: float, interval: str = "") -> int:
    if "hour" in (interval or "").lower() or value < 1000:
        return int(value * 2080)
    return int(value)


def fetch_greenhouse(entry: dict) -> list[Job]:
    data = get_json(f"https://boards-api.greenhouse.io/v1/boards/{entry['token']}/jobs",
                    {"content": "true"})
    jobs = []
    for j in (data or {}).get("jobs", []):
        desc = strip_html(j.get("content", ""))
        loc = (j.get("location") or {}).get("name", "")
        smin = smax = None
        src = ""
        try:
            ranges = j.get("pay_input_ranges") or []
            if ranges:
                smin = _annual(ranges[0]["min_cents"] / 100)
                smax = _annual(ranges[0]["max_cents"] / 100)
                src = "api"
        except (KeyError, TypeError, ValueError):
            smin = smax = None
        if smin is None:
            smin, smax, src = parse_salary(desc)
        jobs.append(Job("greenhouse", entry["company"], j.get("title", ""), loc,
                        j.get("absolute_url", ""), desc, smin, smax, src,
                        remote="remote" in f"{loc} {desc[:400]}".lower()))
    return jobs


def fetch_lever(entry: dict) -> list[Job]:
    data = get_json(f"https://api.lever.co/v0/postings/{entry['slug']}", {"mode": "json"})
    jobs = []
    for j in data or []:
        cats = j.get("categories") or {}
        loc = cats.get("location", "") or ""
        desc = j.get("descriptionPlain") or strip_html(j.get("description", ""))
        for lst in j.get("lists", []) or []:
            desc += "\n" + (lst.get("text", "") or "") + "\n" + strip_html(lst.get("content", ""))
        desc += "\n" + (j.get("additionalPlain") or "")
        smin = smax = None
        src = ""
        sr = j.get("salaryRange")
        if sr and sr.get("min"):
            smin, smax, src = _annual(sr["min"], sr.get("interval", "")), \
                _annual(sr.get("max") or sr["min"], sr.get("interval", "")), "api"
        else:
            smin, smax, src = parse_salary(desc)
        remote = (j.get("workplaceType") == "remote") or "remote" in loc.lower()
        jobs.append(Job("lever", entry["company"], j.get("text", ""), loc,
                        j.get("hostedUrl", ""), desc, smin, smax, src, remote=remote))
    return jobs


def fetch_ashby(entry: dict) -> list[Job]:
    data = get_json(f"https://api.ashbyhq.com/posting-api/job-board/{entry['board']}",
                    {"includeCompensation": "true"})
    jobs = []
    for j in (data or {}).get("jobs", []):
        desc = j.get("descriptionPlain") or strip_html(j.get("descriptionHtml", ""))
        loc = j.get("location", "") or ""
        smin = smax = None
        src = ""
        comp = j.get("compensation") or {}
        for c in comp.get("summaryComponents", []) or []:
            if c.get("compensationType") == "Salary" and c.get("minValue"):
                smin = _annual(c["minValue"], c.get("interval", ""))
                smax = _annual(c.get("maxValue") or c["minValue"], c.get("interval", ""))
                src = "api"
                break
        if smin is None:
            smin, smax, src = parse_salary(comp.get("compensationTierSummary", "") or desc)
        remote = bool(j.get("isRemote")) or "remote" in loc.lower()
        jobs.append(Job("ashby", entry["company"], j.get("title", ""), loc,
                        j.get("jobUrl", ""), desc, smin, smax, src, remote=remote))
    return jobs


def fetch_remotive(cfg_remotive: dict) -> list[Job]:
    jobs = []
    for cat in cfg_remotive.get("categories", []):
        data = get_json("https://remotive.com/api/remote-jobs", {"category": cat})
        for j in (data or {}).get("jobs", []):
            desc = strip_html(j.get("description", ""))
            smin, smax, src = parse_salary(j.get("salary", "") or "") 
            if smin is None:
                smin, smax, src = parse_salary(desc)
            jobs.append(Job("remotive", j.get("company_name", ""), j.get("title", ""),
                            j.get("candidate_required_location", ""), j.get("url", ""),
                            desc, smin, smax, src, remote=True))
        time.sleep(1)  # be polite; Remotive asks for few calls per day
    return jobs


def gather(cfg: dict) -> list[tuple[Job, bool]]:
    """Return (job, trusted) pairs. Curated company boards are 'trusted' (skip the
    education-keyword check); open boards like Remotive are not."""
    src = cfg["sources"]
    out: list[tuple[Job, bool]] = []
    for fetch, key in ((fetch_greenhouse, "greenhouse"), (fetch_lever, "lever"),
                       (fetch_ashby, "ashby")):
        for entry in src.get(key, []) or []:
            try:
                found = fetch(entry)
                log.info("%s/%s: %d postings", key, entry["company"], len(found))
                out += [(j, True) for j in found]
            except Exception as e:  # keep going if one board breaks
                log.warning("%s/%s failed: %s", key, entry.get("company"), e)
            time.sleep(0.5)
    rem = src.get("remotive") or {}
    if rem.get("enabled"):
        try:
            found = fetch_remotive(rem)
            log.info("remotive: %d postings", len(found))
            out += [(j, False) for j in found]
        except Exception as e:
            log.warning("remotive failed: %s", e)
    return out


# --------------------------------------------------------------------------- #
# Seen-jobs memory
# --------------------------------------------------------------------------- #
def open_db(path: str) -> sqlite3.Connection:
    db = sqlite3.connect(path)
    db.execute("CREATE TABLE IF NOT EXISTS seen (k TEXT PRIMARY KEY, first_seen TEXT)")
    return db


def is_seen(db, key: str) -> bool:
    return db.execute("SELECT 1 FROM seen WHERE k=?", (key,)).fetchone() is not None


def mark_seen(db, key: str):
    db.execute("INSERT OR IGNORE INTO seen VALUES (?,?)", (key, dt.date.today().isoformat()))
    db.commit()


# --------------------------------------------------------------------------- #
# LLM steps
# --------------------------------------------------------------------------- #
SCORE_SYSTEM = """You assess job fit for one candidate. Use ONLY the candidate facts provided.
Be realistic and conservative: required skills or years the candidate lacks are gaps.
Return ONLY a JSON object:
{"fit_score": 0-100, "verdict": "apply" | "stretch" | "skip",
 "why": ["up to 3 short reasons it fits"], "gaps": ["requirements she does not clearly meet"]}"""

TAILOR_SYSTEM = """You tailor a resume and cover letter for one candidate and one job.
HARD RULES:
- Use ONLY facts in CANDIDATE FACTS. Never invent employers, tools, metrics, credentials, years, or experience.
- Anything unknown stays as a bracketed placeholder such as [X] or [your LMS platforms].
- If the job requires something the candidate lacks, state that plainly in ONE honest sentence in the
  cover letter and pivot to real strengths. Never hide or fake a gap.
- Never mention health, mental health, harassment, immigration status, gender or orientation.
- Cover letter: 230-300 words, plain and specific, no cliches, addressed to the hiring team.
- Resume summary: at most 60 words. Bullets: 6-9, reordered and reworded for this job from the facts.
Return ONLY a JSON object with keys:
"resume_summary" (string), "resume_bullets" (list of strings), "skills_line" (string),
"cover_letter" (string), "gaps" (list of strings), "questions_for_candidate" (list of strings:
facts you need to strengthen the application)."""


def parse_json(text: str) -> dict:
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("no JSON object in model output")
    return json.loads(text[start:end + 1])


def llm_json(client, model: str, system: str, user: str, max_tokens: int) -> dict:
    msg = client.messages.create(model=model, max_tokens=max_tokens, system=system,
                                 messages=[{"role": "user", "content": user}])
    text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
    return parse_json(text)


def job_block(job: Job) -> str:
    return (f"JOB: {job.title} at {job.company}\nLocation: {job.location}\n"
            f"Salary: {job.salary_display}\nURL: {job.url}\n\n"
            f"DESCRIPTION:\n{job.description[:6000]}")


def score_job(client, cfg: dict, facts: str, job: Job):
    out = llm_json(client, cfg["llm"]["scoring_model"], SCORE_SYSTEM,
                   f"CANDIDATE FACTS:\n{facts}\n\n{job_block(job)}", 700)
    job.fit_score = int(out.get("fit_score", 0))
    job.fit_verdict = out.get("verdict", "")
    job.fit_why = out.get("why", [])[:3]
    job.fit_gaps = out.get("gaps", [])


def tailor_job(client, cfg: dict, facts: str, job: Job) -> dict:
    return llm_json(client, cfg["llm"]["tailoring_model"], TAILOR_SYSTEM,
                    f"CANDIDATE FACTS:\n{facts}\n\n{job_block(job)}", 2500)


# --------------------------------------------------------------------------- #
# Honesty guards on generated text
# --------------------------------------------------------------------------- #
def audit_text(text: str, facts: str, job: Job, forbidden: list[str]) -> list[str]:
    """Flag things a human should double-check before sending."""
    issues = []
    low = text.lower()
    for term in forbidden:
        if re.search(rf"\b{re.escape(term.lower())}\b", low):
            issues.append(f'Mentions "{term}", which is not in your facts. Remove it unless true.')
    known = f"{facts}\n{job.description}\n{job.salary_display}"
    stripped = re.sub(r"\[[^\]]*\]", "", text)          # ignore bracketed placeholders
    for num in set(re.findall(r"\d[\d,\.]*%?", stripped)):
        n = num.rstrip(".,")
        if len(n) > 1 and n not in known:
            issues.append(f"Number {n} is not in your facts or the posting: verify.")
    if re.search(r"\[[^\]]+\]", text):
        issues.append("Contains [placeholders] to fill in.")
    return issues


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #
SUMMARY_HEADERS = ["Date found", "Company", "Role", "Salary", "Fit", "Verdict", "Location",
                   "Link", "Flags", "Application file", "My status"]
POSTING_HEADERS = ["Source", "Company", "Role", "Location", "Salary", "Link", "Status"]
LINK_COL_SUMMARY = (7,)      # zero-based index of the "Link" column
LINK_COL_POSTINGS = (5,)


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:60]


def csv_safe(v) -> str:
    v = "" if v is None else str(v)
    return "'" + v if v[:1] in ("=", "+", "-", "@") else v   # stops spreadsheet formula injection


def write_csv(path: pathlib.Path, headers: list[str], rows: list[list]):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:   # utf-8-sig: Excel-friendly
        w = csv.writer(fh)
        w.writerow(headers)
        for r in rows:
            w.writerow([csv_safe(c) for c in r])


def write_xlsx(path: pathlib.Path, headers: list[str], rows: list[list],
               link_cols=(), append: bool = False) -> pathlib.Path | None:
    """Write (or append to) an .xlsx file. Returns the path written, or None if openpyxl is missing."""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.warning("openpyxl is not installed, so no .xlsx was written (CSV was). "
                    "Run: pip install openpyxl")
        return None

    fresh = True
    if append and path.exists():
        try:
            wb = load_workbook(path)
            ws = wb.active
            fresh = False
        except Exception as e:
            log.warning("could not open existing %s (%s); starting a new file", path.name, e)
    if fresh:
        wb = Workbook()
        ws = wb.active
        ws.title = "Jobs"
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
            c.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"

    for r in rows:
        ws.append(["" if c is None else c for c in r])
        for idx, cell in enumerate(ws[ws.max_row]):
            if isinstance(cell.value, str):
                if cell.value[:1] in ("=", "+", "-", "@"):
                    cell.data_type = "s"                    # treat as text, never a formula
                elif idx in link_cols and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
    if fresh:
        for i, h in enumerate(headers, start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)] + [8])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    ws.auto_filter.ref = ws.dimensions

    try:
        wb.save(path)
        return path
    except PermissionError:      # file is open in Excel
        alt = path.with_name(f"{path.stem}_new{path.suffix}")
        wb.save(alt)
        log.warning("%s is open in another program; saved %s instead", path.name, alt.name)
        return alt


def job_rows(jobs: list[Job], files: dict[str, pathlib.Path], date: str) -> list[list]:
    rows = []
    for j in sorted(jobs, key=lambda j: j.fit_score if j.fit_score is not None else -1, reverse=True):
        f = files.get(j.key)
        rows.append([date, j.company, j.title, j.salary_display,
                     "" if j.fit_score is None else j.fit_score, j.fit_verdict, j.location,
                     j.url, "; ".join(j.flags), f.name if f else "", ""])
    return rows


def update_tracker(base: pathlib.Path, rows: list[list]) -> list[pathlib.Path]:
    """Cumulative list of every match across days. Appends, so your 'My status' edits survive."""
    written = []
    csv_path = base / "job_tracker.csv"
    new_file = not csv_path.exists()
    with open(csv_path, "a", newline="", encoding="utf-8-sig" if new_file else "utf-8") as fh:
        w = csv.writer(fh)
        if new_file:
            w.writerow(SUMMARY_HEADERS)
        for r in rows:
            w.writerow([csv_safe(c) for c in r])
    written.append(csv_path)
    x = write_xlsx(base / "job_tracker.xlsx", SUMMARY_HEADERS, rows, LINK_COL_SUMMARY, append=True)
    if x:
        written.append(x)
    return written


def write_reports(out_dir: pathlib.Path, base: pathlib.Path, stamp: str, date: str,
                  jobs: list[Job], files: dict[str, pathlib.Path], stats: dict,
                  all_rows: list[list], notes: list[str], prefix: str = "summary",
                  track: bool = True) -> list[pathlib.Path]:
    """Always writes a report, even when zero jobs matched, so a run is never silent."""
    rows = job_rows(jobs, files, date)
    md = [f"# Remote job matches: {stamp}", ""]
    md += [f"> {n}" for n in notes]
    md += ["", f"Fetched **{stats['fetched']}** postings. **{len(jobs)}** new matches.", ""]
    if stats["rejected"]:
        md += ["Why other postings were filtered out:", ""]
        md += [f"- {k}: {v}" for k, v in sorted(stats["rejected"].items(), key=lambda kv: -kv[1])]
        md += [""]
    if rows:
        md += ["| Company | Role | Salary | Fit | Flags | Application |", "|---|---|---|---|---|---|"]
        for r in rows:
            link = f"[{r[9]}](jobs/{r[9]})" if r[9] else "not tailored"
            md.append(f"| {r[1]} | [{r[2]}]({r[7]}) | {r[3]} | {r[4] if r[4] != '' else 'n/a'} | "
                      f"{r[8]} | {link} |")
    else:
        md.append("No new matching jobs this run.")

    written: list[pathlib.Path] = []
    p = out_dir / f"{prefix}_{stamp}.md"
    p.write_text("\n".join(md) + "\n", encoding="utf-8")
    written.append(p)
    p = out_dir / f"{prefix}_{stamp}.csv"
    write_csv(p, SUMMARY_HEADERS, rows)
    written.append(p)
    x = write_xlsx(out_dir / f"{prefix}_{stamp}.xlsx", SUMMARY_HEADERS, rows, LINK_COL_SUMMARY)
    if x:
        written.append(x)
    p = out_dir / f"all_postings_{stamp}.csv"          # every posting seen + why it was kept/dropped
    write_csv(p, POSTING_HEADERS, all_rows)
    written.append(p)
    if track:
        done = [j for j in jobs if not any("failed" in f.lower() or "not scored" in f.lower()
                                           for f in j.flags)]
        if done:
            written += update_tracker(base, job_rows(done, files, date))
    return written


def write_job_file(folder: pathlib.Path, job: Job, t: dict, issues: list[str]) -> pathlib.Path:
    lines = [
        f"# {job.title} at {job.company}", "",
        f"- **Salary:** {job.salary_display}",
        f"- **Location:** {job.location}",
        f"- **Link:** {job.url}",
        f"- **Fit:** {job.fit_score}/100 ({job.fit_verdict})", "",
    ]
    if job.flags:
        lines += ["**Flags:** " + "; ".join(job.flags), ""]
    if job.fit_gaps or t.get("gaps"):
        lines += ["## Gaps (be ready to address)"] + [f"- {g}" for g in (t.get("gaps") or job.fit_gaps)] + [""]
    if issues:
        lines += ["## Review before sending"] + [f"- {i}" for i in issues] + [""]
    if t.get("questions_for_candidate"):
        lines += ["## Facts that would strengthen this"] + [f"- {q}" for q in t["questions_for_candidate"]] + [""]
    lines += ["## Tailored resume", "", f"**Summary:** {t.get('resume_summary', '')}", "",
              "**Klett World Languages, EdTech Support Specialist (lead bullets)**", ""]
    lines += [f"- {b}" for b in t.get("resume_bullets", [])]
    lines += ["", f"**Skills:** {t.get('skills_line', '')}", "", "## Cover letter", "",
              t.get("cover_letter", ""), ""]
    path = folder / f"{slug(job.company)}__{slug(job.title)}.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def load_cfg(path: str) -> dict:
    with open(path, encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def unmark_seen(db, key: str):
    db.execute("DELETE FROM seen WHERE k=?", (key,))
    db.commit()


def run(cfg_path: str, dry_run: bool, limit: int | None, allow_unknown_salary: bool = False):
    cfg = load_cfg(cfg_path)
    if allow_unknown_salary:
        cfg["profile"]["allow_unknown_salary"] = True
    base = pathlib.Path(cfg_path).resolve().parent
    facts = (base / cfg["profile"]["resume_facts"]).read_text(encoding="utf-8")
    now = dt.datetime.now()
    today, stamp = now.date().isoformat(), now.strftime("%Y-%m-%d_%H%M%S")
    out_dir = base / cfg.get("output_dir", "output") / today
    (out_dir / "jobs").mkdir(parents=True, exist_ok=True)
    db = open_db(str(base / "seen.db"))
    prefix = "dryrun" if dry_run else "summary"

    # ---- 1. fetch + filter ------------------------------------------------- #
    candidates: list[Job] = []
    rejected: dict[str, int] = {}
    all_rows: list[list] = []
    in_run: set[str] = set()
    fetched = 0
    for job, trusted in gather(cfg):
        if job.key in in_run:            # same posting returned by two feeds/categories
            continue
        in_run.add(job.key)
        fetched += 1
        ok, why = apply_filters(job, cfg, trusted)
        status = "kept" if ok else f"filtered: {why}"
        if ok and is_seen(db, job.key):
            ok, why, status = False, "already seen", "filtered: already seen"
        if ok:
            candidates.append(job)
        else:
            rejected[why] = rejected.get(why, 0) + 1
        all_rows.append([job.source, job.company, job.title, job.location,
                         job.salary_display, job.url, status])
    log.info("fetched %d; kept %d new; filtered: %s", fetched, len(candidates), rejected)
    if limit:
        candidates = candidates[:limit]
    stats = {"fetched": fetched, "rejected": rejected}

    notes: list[str] = []
    if fetched == 0:
        notes.append("No postings were fetched at all. Check your internet connection, the board "
                     "slugs in config.yaml, and the log lines above.")
    elif not candidates and rejected.get("salary unknown", 0) > fetched * 0.5:
        notes.append("Most postings were dropped because no salary was listed. To include them, "
                     "run with --allow-unknown-salary (or set allow_unknown_salary: true).")

    def finish(jobs, files, track):
        paths = write_reports(out_dir, base, stamp, today, jobs, files, stats, all_rows,
                              notes, prefix=prefix, track=track)
        for n in notes:
            print(f"- {n}")
        print(f"\n{fetched} postings fetched, {len(jobs)} new matches.")
        print("Files written:")
        for p in paths:
            print(f"  {p}")

    # ---- 2. nothing to do, or dry run ------------------------------------- #
    if dry_run:
        notes.append("Dry run: nothing was scored, tailored, or marked as seen.")
        finish(candidates, {}, track=False)
        return
    if not candidates:
        finish([], {}, track=False)
        return

    # ---- 3. start the LLM client (fall back to an unscored report) --------- #
    llm = cfg["llm"]
    client = None
    if not os.environ.get("ANTHROPIC_API_KEY"):
        notes.append("ANTHROPIC_API_KEY is not set, so jobs were NOT scored or tailored. "
                     "Set it and rerun; these jobs will be picked up again.")
    else:
        try:
            from anthropic import Anthropic
            client = Anthropic()
        except Exception as e:
            notes.append(f"Could not start the Anthropic client ({e}). Run: pip install anthropic. "
                         "Jobs were NOT scored or tailored, and will be picked up again.")
    if client is None:
        for j in candidates:
            j.flags.append("Not scored")
        finish(candidates, {}, track=False)       # not marked as seen, so tomorrow retries them
        return

    # ---- 4. score, then tailor the best few -------------------------------- #
    scored: list[Job] = []
    cap = llm.get("max_scored_per_day", 40)
    for job in candidates[:cap]:
        try:
            score_job(client, cfg, facts, job)
            scored.append(job)
            mark_seen(db, job.key)               # only after a successful score
        except Exception as e:
            log.warning("scoring failed for %s: %s", job.title, e)
            job.flags.append("Scoring failed: will retry next run")
            scored.append(job)
    if len(candidates) > cap:
        notes.append(f"{len(candidates) - cap} more new jobs were held for tomorrow (daily cap of {cap}).")

    shortlist = sorted((j for j in scored if (j.fit_score or 0) >= llm.get("min_fit", 60)
                        and j.fit_verdict != "skip"),
                       key=lambda j: j.fit_score or 0, reverse=True)[: llm.get("max_tailored_per_day", 3)]

    files: dict[str, pathlib.Path] = {}
    forbidden = cfg["profile"].get("forbidden_claims", [])
    for job in shortlist:
        try:
            t = tailor_job(client, cfg, facts, job)
            body = " ".join([t.get("resume_summary", ""), " ".join(t.get("resume_bullets", [])),
                             t.get("skills_line", ""), t.get("cover_letter", "")])
            issues = audit_text(body, facts, job, forbidden)
            files[job.key] = write_job_file(out_dir / "jobs", job, t, issues)
        except Exception as e:
            log.warning("tailoring failed for %s: %s", job.title, e)
            job.flags.append("Tailoring failed: will retry next run")
            unmark_seen(db, job.key)

    finish(scored, files, track=True)


# --------------------------------------------------------------------------- #
# Offline self-test
# --------------------------------------------------------------------------- #
def selftest() -> int:
    fails = 0

    def check(name, got, want):
        nonlocal fails
        ok = got == want
        fails += not ok
        print(("PASS " if ok else "FAIL ") + name + ("" if ok else f" -> got {got!r}, want {want!r}"))

    ps = lambda s: parse_salary(s)[:2]
    check("range with $", ps("The salary range is $80,000 - $95,000 USD"), (80000, 95000))
    check("k range", ps("Compensation: $70K–$80K per year"), (70000, 80000))
    check("USD prefix", ps("Salary USD 90,000 - 130,000"), (90000, 130000))
    check("hourly single", ps("Pay: $35/hour"), (72800, 72800))
    check("ignores funding + 401k", ps("We raised $35M and offer a 401(k) match."), (None, None))
    check("ignores CA$", ps("CA$85K - CA$106K"), (None, None))
    check("khan style", ps("target salary range is $96,800 $119,340 USD $130,807 $161,109 CAN"), (96800, 119340))
    check("below floor still parses", ps("$24.00 per hour"), (49920, 49920))

    cfg = {"profile": {"home_state": "California", "min_salary": 70000, "salary_basis": "max",
                       "max_travel_percent": 25, "allow_unknown_salary": False},
           "filters": {"title_include": [r"support", r"customer success", r"implementation", r"integration"],
                       "title_exclude": [r"director", r"\bvp\b", r"software engineer", r"sales"],
                       "exclude_location_tokens": ["canada", "india", "united kingdom", "mexico"],
                       "domain_keywords": ["k-12", "education"]}}
    check("loc: US remote", location_eligibility("Remote - US", cfg), "ok")
    check("loc: California", location_eligibility("Remote - California", cfg), "ok")
    check("loc: other state only", location_eligibility("Remote - Texas", cfg), "no")
    check("loc: India", location_eligibility("India (Remote)", cfg), "no")
    check("loc: US or Canada", location_eligibility("Remote - US or Canada", cfg), "ok")
    check("loc: bare remote", location_eligibility("Remote", cfg), "check")
    check("travel 65", max_travel_percent("Travel up to 65% required"), 65)
    check("travel 15-50", max_travel_percent("may require approximately 15-50% travel"), 50)
    check("no travel", max_travel_percent("Fully remote."), None)
    check("title ok", title_ok("Technical Support Specialist", cfg), True)
    check("title excluded", title_ok("Director of Customer Success", cfg), False)

    def mk(**kw):
        d = dict(source="t", company="Acme", title="Customer Success Manager", location="Remote - US",
                 url="u", description="K-12 education platform.", salary_min=80000, salary_max=95000,
                 salary_source="api", remote=True)
        d.update(kw)
        return Job(**d)

    check("filter pass", apply_filters(mk(), cfg, True)[0], True)
    check("filter: low salary", apply_filters(mk(salary_min=50000, salary_max=60000), cfg, True)[1], "salary below floor")
    check("filter: unknown salary", apply_filters(mk(salary_min=None, salary_max=None), cfg, True)[1], "salary unknown")
    check("filter: travel", apply_filters(mk(description="Travel up to 65%"), cfg, True)[1], "travel 65%")
    check("filter: not remote", apply_filters(mk(remote=False), cfg, True)[1], "not remote")
    check("filter: untrusted off-domain", apply_filters(mk(description="Dental billing"), cfg, False)[1],
          "not education/nonprofit")

    facts = "Closed 536 tickets. 65% first-contact resolution. Klett 2.5 years."
    issues = audit_text("I know Python and closed 900 tickets with 65% FCR. [X] partners.", facts,
                        mk(), ["Python", "dbt"])
    check("audit flags Python", any("Python" in i for i in issues), True)
    check("audit flags unknown number", any("900" in i for i in issues), True)
    check("audit allows known number", any("65%" in i for i in issues), False)
    check("audit notes placeholders", any("placeholder" in i for i in issues), True)

    check("parse_json w/ prose", parse_json('Sure:\n{"a": 1}\nDone'), {"a": 1})

    # ---- report files must be written, even with zero matches --------------- #
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        base = pathlib.Path(tmp)
        out = base / "output"
        out.mkdir()
        stats = {"fetched": 3, "rejected": {"salary unknown": 3}}
        paths = write_reports(out, base, "2026-01-01_0900", "2026-01-01", [], {}, stats, [], [])
        names = {p.suffix for p in paths}
        check("zero matches still writes md/csv/xlsx", {".md", ".csv", ".xlsx"} <= names, True)
        check("zero matches: files exist", all(p.exists() for p in paths), True)
        j = mk(fit_score=80, fit_verdict="apply", title="=HYPERLINK(\"x\")", url="https://x/1")
        allrows = [["t", "Acme", j.title, "Remote", j.salary_display, j.url, "kept"]]
        paths = write_reports(out, base, "2026-01-01_0901", "2026-01-01", [j], {}, stats, allrows, [])
        rows = list(csv.reader(open(out / "summary_2026-01-01_0901.csv", encoding="utf-8-sig")))
        check("csv has header + 1 job", len(rows), 2)
        check("csv neutralizes formulas", rows[1][2].startswith("'="), True)
        write_reports(out, base, "2026-01-02_0900", "2026-01-02", [mk(fit_score=70, url="https://x/2")],
                      {}, stats, allrows, [])
        trk = list(csv.reader(open(base / "job_tracker.csv", encoding="utf-8-sig")))
        check("tracker accumulates across runs", len(trk), 3)
        try:
            from openpyxl import load_workbook
            ws = load_workbook(base / "job_tracker.xlsx").active
            check("xlsx tracker has header + 2 rows", ws.max_row, 3)
            check("xlsx link is clickable", ws.cell(2, 8).hyperlink is not None, True)
            check("xlsx formula stored as text", ws.cell(2, 3).data_type, "s")
        except ImportError:
            print("SKIP xlsx checks (openpyxl not installed)")
    print(f"\n{'ALL TESTS PASSED' if not fails else f'{fails} TEST(S) FAILED'}")
    return 1 if fails else 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default=str(HERE / "config.yaml"))
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--limit", type=int)
    ap.add_argument("--allow-unknown-salary", action="store_true",
                    help="include postings that don't list a salary")
    ap.add_argument("--selftest", action="store_true")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    if args.selftest:
        sys.exit(selftest())
    run(args.config, args.dry_run, args.limit, args.allow_unknown_salary)
